from pptx import Presentation
from pptx.util import Inches, Pt
import json

prs = Presentation(r'c:/Users/atanu/Desktop/TEST/Sample.pptx')
print(f'Slides: {len(prs.slides)}')
print(f'Slide size: {prs.slide_width.inches:.2f}" x {prs.slide_height.inches:.2f}"')
print(f'Layouts available: {[l.name for l in prs.slide_layouts]}')

for i, sl in enumerate(prs.slides):
    layout = sl.slide_layout.name
    print(f'\n--- Slide {i+1} [layout={layout}] ---')
    for s in sl.shapes:
        txt = ""
        if hasattr(s, "text"):
            txt = s.text[:80].replace("\n", " | ")
        print(f'  shape_type={s.shape_type} name={s.name!r} left={s.left} top={s.top} w={s.width} h={s.height} text={txt!r}')
